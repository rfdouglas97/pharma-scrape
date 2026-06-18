import io

from openpyxl import Workbook

from pipeline_intel.ingest.parse_doc import (
    doc_kind,
    parse_document,
    render_markdown_table,
)


def test_doc_kind_from_ext_and_content_type():
    assert doc_kind(ext=".csv") == "csv"
    assert doc_kind(ext=".xlsx") == "xlsx"
    assert doc_kind(ext=".xls") == "xlsx"
    assert doc_kind(ext=".pdf") == "pdf"
    assert doc_kind(content_type="application/pdf") == "pdf"
    assert doc_kind(content_type="text/csv; charset=utf-8") == "csv"
    assert doc_kind(content_type="application/octet-stream", ext=".xlsx") == "xlsx"
    assert doc_kind(content_type="text/html") is None
    assert doc_kind() is None


def test_render_markdown_table_uses_first_row_as_header():
    md = render_markdown_table([["Asset", "Phase"], ["ABC-1", "Phase 2"]])
    lines = md.splitlines()
    assert lines[0] == "| Asset | Phase |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| ABC-1 | Phase 2 |"


def test_render_markdown_table_pads_ragged_rows_and_skips_blank():
    md = render_markdown_table([["A", "B", "C"], [], ["x"]])
    lines = md.splitlines()
    assert lines[0] == "| A | B | C |"
    assert lines[-1] == "| x |  |  |"


def test_parse_csv_renders_table_and_keeps_rows():
    raw = b"Asset,Indication,Phase\nABC-1,NSCLC,Phase 2\nXYZ-9,Melanoma,Phase 1\n"
    doc = parse_document(raw, ext=".csv")
    assert doc.kind == "csv"
    assert "| Asset | Indication | Phase |" in doc.text
    assert "| ABC-1 | NSCLC | Phase 2 |" in doc.text
    assert doc.tables[0][2] == ["XYZ-9", "Melanoma", "Phase 1"]


def test_parse_csv_sniffs_semicolon_delimiter():
    raw = b"Asset;Phase\nABC-1;Phase 3\n"
    doc = parse_document(raw, content_type="text/csv")
    assert "| Asset | Phase |" in doc.text
    assert "| ABC-1 | Phase 3 |" in doc.text


def test_parse_xlsx_renders_each_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    ws.append(["Asset", "Indication", "Phase"])
    ws.append(["ABC-1", "NSCLC", "Phase 2"])
    buf = io.BytesIO()
    wb.save(buf)

    doc = parse_document(buf.getvalue(), ext=".xlsx")
    assert doc.kind == "xlsx"
    assert "## Sheet: Pipeline" in doc.text
    assert "| Asset | Indication | Phase |" in doc.text
    assert "| ABC-1 | NSCLC | Phase 2 |" in doc.text
    assert doc.tables[0][1] == ["ABC-1", "NSCLC", "Phase 2"]


def test_parse_document_rejects_unknown_type():
    try:
        parse_document(b"<html></html>", content_type="text/html")
    except ValueError as exc:
        assert "unsupported document type" in str(exc)
    else:
        raise AssertionError("expected ValueError for unsupported type")
